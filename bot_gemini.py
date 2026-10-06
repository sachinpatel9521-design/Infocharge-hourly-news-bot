# INFOCHARGE FINAL BUILD: V4.1-2026-10-06 — syntax-validated
import os
import json
import logging
import re
import html
import time
import math
import xml.etree.ElementTree as ET
from pathlib import Path
from datetime import datetime, timedelta
from zoneinfo import ZoneInfo

import requests
from google import genai
from google.genai import types

try:
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
except Exception:
    plt = None

# ============================================================
# INFOCHARGE MARKET INTELLIGENCE BOT — V4
# Free-first architecture:
#   RSS discovery + public market/calendar data + Gemini reasoning
#   No paid Google Search grounding required.
#
# Features:
#   1. Very selective Indian stock/sector intelligence
#   2. Geopolitics -> commodity/currency/trade -> Indian stock/sector link
#   3. Fed/FOMC, US CPI, inflation, Treasury yields, RBI/SEBI/NSE events
#   4. Educational market-learning posts and historical thinking frameworks
#   5. Nifty IT / Nasdaq ratio educational snapshot when data is available
#   6. IPO watch / due-diligence posts from discovered public information
#   7. Occasional market-related motivation, never spammy
#   8. Persistent editorial memory
#   9. Automatic branded visual cards (free; no image-generation API)
#  10. Gemini retry + 3.7 fallback for temporary 503s
#
# IMPORTANT:
#   This is educational market intelligence, not investment advice.
# ============================================================

IST = ZoneInfo("Asia/Kolkata")

TELEGRAM_BOT_TOKEN = os.environ["TELEGRAM_BOT_TOKEN"]
TELEGRAM_CHAT_ID = os.environ["TELEGRAM_CHAT_ID"]
GEMINI_API_KEY = os.environ["GEMINI_API_KEY"]

PRIMARY_MODEL = "gemini-3.8-flash"
FALLBACK_MODEL = "gemini-3.7-flash"

MEMORY_FILE = Path("infocharge_memory.json")
MAX_MEMORY_ITEMS = 140
MAX_ARTICLES = 90
HTTP_TIMEOUT = 18
GEMINI_RETRIES = 4
MIN_IMPORTANCE = 8

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s | %(levelname)s | %(message)s"
)

client = genai.Client(api_key=GEMINI_API_KEY)

# ============================================================
# PUBLIC DISCOVERY SOURCES
# ============================================================

NEWS_FEEDS = [
    # Broad but market-focused
    "https://news.google.com/rss/search?q=India+listed+company+order+contract+capex+results&hl=en-IN&gl=IN&ceid=IN:en",
    "https://news.google.com/rss/search?q=India+NSE+BSE+corporate+announcement+stocks&hl=en-IN&gl=IN&ceid=IN:en",
    "https://news.google.com/rss/search?q=India+SEBI+RBI+market+policy+stocks&hl=en-IN&gl=IN&ceid=IN:en",

    # Sector intelligence
    "https://news.google.com/rss/search?q=India+data+centre+AI+infrastructure+stocks+order&hl=en-IN&gl=IN&ceid=IN:en",
    "https://news.google.com/rss/search?q=India+semiconductor+electronics+manufacturing+stocks+investment&hl=en-IN&gl=IN&ceid=IN:en",
    "https://news.google.com/rss/search?q=India+defence+order+contract+listed+company&hl=en-IN&gl=IN&ceid=IN:en",
    "https://news.google.com/rss/search?q=India+pharma+USFDA+EMA+approval+listed+company&hl=en-IN&gl=IN&ceid=IN:en",
    "https://news.google.com/rss/search?q=India+power+renewable+grid+order+listed+company&hl=en-IN&gl=IN&ceid=IN:en",
    "https://news.google.com/rss/search?q=India+bank+NPA+credit+deposit+RBI+listed+bank&hl=en-IN&gl=IN&ceid=IN:en",
    "https://news.google.com/rss/search?q=India+IT+services+deal+contract+guidance+listed+company&hl=en-IN&gl=IN&ceid=IN:en",
    "https://news.google.com/rss/search?q=India+chemical+specialty+chemical+export+import+stocks&hl=en-IN&gl=IN&ceid=IN:en",
    "https://news.google.com/rss/search?q=India+cement+steel+commodity+capex+infrastructure+stocks&hl=en-IN&gl=IN&ceid=IN:en",

    # Geopolitical transmission channels
    "https://news.google.com/rss/search?q=India+geopolitics+tariff+trade+sanctions+stocks+sector&hl=en-IN&gl=IN&ceid=IN:en",
    "https://news.google.com/rss/search?q=Middle+East+oil+shipping+India+stocks+sector&hl=en-IN&gl=IN&ceid=IN:en",
    "https://news.google.com/rss/search?q=Russia+Ukraine+India+trade+energy+stocks&hl=en-IN&gl=IN&ceid=IN:en",
    "https://news.google.com/rss/search?q=US+China+trade+tariff+India+stocks+sector&hl=en-IN&gl=IN&ceid=IN:en",
    "https://news.google.com/rss/search?q=India+China+border+trade+industry+stocks&hl=en-IN&gl=IN&ceid=IN:en",
    "https://news.google.com/rss/search?q=Red+Sea+shipping+freight+India+stocks&hl=en-IN&gl=IN&ceid=IN:en",

    # IPO
    "https://news.google.com/rss/search?q=India+IPO+DRHP+RHP+price+band+listing+2026&hl=en-IN&gl=IN&ceid=IN:en",
    "https://news.google.com/rss/search?q=upcoming+IPO+India+2026+NSE+BSE&hl=en-IN&gl=IN&ceid=IN:en",

    # Macro / market structure
    "https://news.google.com/rss/search?q=US+Fed+CPI+inflation+Treasury+yield+India+market&hl=en-IN&gl=IN&ceid=IN:en",
    "https://news.google.com/rss/search?q=US+10+year+yield+India+equity+market+IT+stocks&hl=en-IN&gl=IN&ceid=IN:en",
    "https://news.google.com/rss/search?q=India+rupee+crude+oil+FII+equity+market&hl=en-IN&gl=IN&ceid=IN:en",
]

# Official-ish calendar/release URLs used as deterministic reminders.
FOMC_URL = "https://www.federalreserve.gov/monetarypolicy/fomccalendars.htm"
CPI_URL = "https://www.bls.gov/schedule/news_release/cpi.htm"
NSE_HOLIDAY_URL = "https://www.nseindia.com/resources/exchange-communication-holidays"

NSE_HOLIDAYS_2026 = {
    "2026-01-15", "2026-01-26", "2026-03-03", "2026-03-26",
    "2026-03-31", "2026-04-03", "2026-04-14", "2026-05-01",
    "2026-05-28", "2026-06-26", "2026-09-14", "2026-10-02",
    "2026-10-20", "2026-11-10", "2026-11-24", "2026-12-25",
    # NSE says Nov 8 is a trading holiday; Muhurat Trading is scheduled separately.
    "2026-11-08",
}

# FOMC 2026 dates. Keep explicit so reminders are deterministic.
FOMC_DATES_2026 = {
    "2026-01-27", "2026-01-28",
    "2026-03-17", "2026-03-18",
    "2026-04-28", "2026-04-29",
    "2026-06-16", "2026-06-17",
    # Add remaining dates if/when official calendar is updated.
    "2026-07-28", "2026-07-29",
    "2026-09-15", "2026-09-16",
    "2026-10-27", "2026-10-28",
    "2026-12-08", "2026-12-09",
}

# 2026 BLS CPI release dates from the official BLS calendar.
CPI_RELEASE_DATES_2026 = {
    "2026-01-13", "2026-02-13", "2026-03-11", "2026-04-10",
    "2026-05-12", "2026-06-10", "2026-07-14", "2026-08-12",
    "2026-09-11", "2026-10-14", "2026-11-10", "2026-12-10",
}

# ============================================================
# BASIC HELPERS
# ============================================================

def clean_text(value):
    value = html.unescape(value or "")
    value = re.sub(r"<[^>]+>", " ", value)
    return re.sub(r"\s+", " ", value).strip()


def normalize_title(title):
    title = clean_text(title).lower()
    title = re.sub(r"[^a-z0-9]+", " ", title)
    return re.sub(r"\s+", " ", title).strip()


def safe_int(value, default=0):
    try:
        return int(float(value))
    except Exception:
        return default


def load_memory():
    if not MEMORY_FILE.exists():
        return {
            "posted": [],
            "editorial_notes": [],
            "education": [],
            "ipo_watch": [],
            "calendar_sent": [],
        }

    try:
        data = json.loads(MEMORY_FILE.read_text(encoding="utf-8"))
        if not isinstance(data, dict):
            raise ValueError("Memory is not an object")

        for key in ("posted", "editorial_notes", "education", "ipo_watch", "calendar_sent"):
            data.setdefault(key, [])

        return data

    except Exception as exc:
        logging.warning("Memory read failed: %s", exc)
        return {
            "posted": [],
            "editorial_notes": [],
            "education": [],
            "ipo_watch": [],
            "calendar_sent": [],
        }


def save_memory(memory):
    MEMORY_FILE.write_text(
        json.dumps(memory, ensure_ascii=False, indent=2),
        encoding="utf-8"
    )


def memory_for_prompt(memory):
    compact = {
        "recent_posts": [
            {
                "date": x.get("date"),
                "headline": x.get("headline"),
                "company": x.get("company"),
                "key": x.get("key"),
            }
            for x in memory.get("posted", [])[-MAX_MEMORY_ITEMS:]
        ],
        "editorial_notes": memory.get("editorial_notes", [])[-30:],
        "recent_education": memory.get("education", [])[-15:],
        "ipo_watch": memory.get("ipo_watch", [])[-20:],
    }
    return json.dumps(compact, ensure_ascii=False)


def market_status(now):
    date_key = now.strftime("%Y-%m-%d")

    if now.weekday() >= 5 or date_key in NSE_HOLIDAYS_2026:
        return "closed"

    minutes = now.hour * 60 + now.minute

    if minutes < 555:
        return "pre_market"
    if minutes <= 930:
        return "market_hours"
    return "post_market"


# ============================================================
# NEWS COLLECTION
# ============================================================

def fetch_news():
    collected = []
    seen = set()

    for feed_url in NEWS_FEEDS:
        try:
            response = requests.get(
                feed_url,
                timeout=HTTP_TIMEOUT,
                headers={"User-Agent": "INFOCHARGE-Market-Intelligence/4.0"},
            )
            response.raise_for_status()

            root = ET.fromstring(response.content)

            for item in root.findall(".//item")[:15]:

                def get_child(tag):
                    node = item.find(tag)
                    return node.text if node is not None and node.text else ""

                title = clean_text(get_child("title"))
                summary = clean_text(get_child("description"))
                link = clean_text(get_child("link"))
                published = clean_text(get_child("pubDate"))

                if not title:
                    continue

                key = normalize_title(title)

                if not key or key in seen:
                    continue

                seen.add(key)

                collected.append({
                    "title": title,
                    "summary": summary[:1200],
                    "published": published,
                    "link": link,
                })

        except Exception as exc:
            logging.warning("News feed failed: %s", exc)

    return collected[:MAX_ARTICLES]


def article_packet(articles):
    blocks = []

    for i, item in enumerate(articles, 1):
        blocks.append(
            "\n".join([
                f"ARTICLE {i}",
                f"Published: {item['published']}",
                f"Title: {item['title']}",
                f"Summary: {item['summary']}",
                f"Link: {item['link']}",
            ])
        )

    return "\n\n".join(blocks)


# ============================================================
# GEMINI WITH RETRIES / FALLBACK
# ============================================================

def call_gemini(prompt, max_tokens):
    models = [PRIMARY_MODEL, FALLBACK_MODEL]
    last_error = None

    for model in models:
        for attempt in range(GEMINI_RETRIES):
            try:
                logging.info(
                    "Gemini request | model=%s | attempt=%s",
                    model,
                    attempt + 1
                )

                response = client.models.generate_content(
                    model=model,
                    contents=prompt,
                    config=types.GenerateContentConfig(
                        max_output_tokens=max_tokens
                    ),
                )

                text = (response.text or "").strip()

                if not text:
                    raise RuntimeError("Gemini returned an empty response.")

                return text

            except Exception as exc:
                last_error = exc
                logging.warning(
                    "Gemini failed | model=%s | attempt=%s | %s",
                    model,
                    attempt + 1,
                    exc
                )

                # Retry temporary overload/server/rate-limit failures.
                error_text = str(exc).upper()
                retryable = any(
                    token in error_text
                    for token in (
                        "503", "UNAVAILABLE", "429", "RESOURCE_EXHAUSTED",
                        "500", "INTERNAL", "DEADLINE", "TIMEOUT"
                    )
                )

                if not retryable:
                    break

                time.sleep(3 * (2 ** attempt))

    raise RuntimeError(f"Gemini failed after retries: {last_error}")


def extract_json(text):
    text = text.strip()
    text = re.sub(r"^```(?:json)?\s*", "", text, flags=re.IGNORECASE)
    text = re.sub(r"\s*```$", "", text)

    try:
        return json.loads(text)
    except json.JSONDecodeError:
        match = re.search(r"\{.*\}", text, flags=re.DOTALL)

        if not match:
            raise RuntimeError("Gemini did not return valid JSON.")

        return json.loads(match.group(0))


# ============================================================
# MARKET DATA
# ============================================================

def yahoo_chart(symbol, period1=None, period2=None, interval="1d"):
    """
    Public Yahoo chart endpoint. Used only for educational snapshots.
    No API key is stored.
    """
    url = f"https://query1.finance.yahoo.com/v8/finance/chart/{symbol}"

    params = {
        "interval": interval,
        "range": "1y" if not period1 else None,
    }

    if period1:
        params["period1"] = int(period1.timestamp())

    if period2:
        params["period2"] = int(period2.timestamp())

    try:
        response = requests.get(
            url,
            params={k: v for k, v in params.items() if v is not None},
            timeout=HTTP_TIMEOUT,
            headers={"User-Agent": "Mozilla/5.0"},
        )
        response.raise_for_status()
        return response.json()
    except Exception as exc:
        logging.warning("Yahoo chart failed for %s: %s", symbol, exc)
        return None


def latest_yahoo_price(symbol):
    data = yahoo_chart(symbol, interval="1d")

    try:
        result = data["chart"]["result"][0]
        meta = result.get("meta", {})
        price = meta.get("regularMarketPrice")

        if price is not None:
            return float(price)

        closes = result["indicators"]["quote"][0]["close"]
        closes = [x for x in closes if x is not None]

        return float(closes[-1]) if closes else None

    except Exception:
        return None


def nifty_it_nasdaq_ratio():
    """
    Educational ratio:
      NIFTY IT index proxy / NASDAQ-100 proxy.

    Yahoo symbols can change. If either series is unavailable,
    the bot simply skips the ratio rather than inventing data.
    """
    nifty_it_candidates = ["^CNXIT", "^NIFTYIT"]
    nasdaq_candidates = ["^NDX"]

    it_price = None
    ndx_price = None
    it_symbol = None
    ndx_symbol = None

    for symbol in nifty_it_candidates:
        it_price = latest_yahoo_price(symbol)
        if it_price is not None:
            it_symbol = symbol
            break

    for symbol in nasdaq_candidates:
        ndx_price = latest_yahoo_price(symbol)
        if ndx_price is not None:
            ndx_symbol = symbol
            break

    if it_price is None or ndx_price is None or ndx_price == 0:
        return None

    return {
        "nifty_it": it_price,
        "nasdaq100": ndx_price,
        "ratio": it_price / ndx_price,
        "symbols": [it_symbol, ndx_symbol],
    }


def treasury_10y():
    """
    Uses the U.S. Treasury daily CSV endpoint.
    """
    url = (
        "https://home.treasury.gov/resource-center/data-chart-center/"
        "interest-rates/daily-treasury-rates.csv/2026/all"
        "?type=daily_treasury_yield_curve"
    )

    try:
        response = requests.get(
            url,
            timeout=HTTP_TIMEOUT,
            headers={"User-Agent": "INFOCHARGE-Market-Intelligence/4.0"},
        )
        response.raise_for_status()

        lines = response.text.splitlines()

        if not lines:
            return None

        header = [x.strip().replace("\ufeff", "") for x in lines[0].split(",")]

        # Treasury CSV occasionally contains different column names.
        ten_candidates = [
            "10 YR",
            "10-year",
            "10-year Treasury",
            "10 Yr",
        ]

        idx = None

        for candidate in ten_candidates:
            if candidate in header:
                idx = header.index(candidate)
                break

        if idx is None:
            for i, h in enumerate(header):
                if "10" in h and "YR" in h.upper():
                    idx = i
                    break

        if idx is None:
            return None

        for row in reversed(lines[1:]):
            parts = row.split(",")

            if len(parts) <= idx:
                continue

            date_value = parts[0].strip()
            value = parts[idx].strip()

            if value and value != "N/A":
                return {
                    "date": date_value,
                    "yield": float(value),
                }

    except Exception as exc:
        logging.warning("Treasury 10Y fetch failed: %s", exc)

    return None



def fred_latest(series_id):
    """Read a public FRED CSV without requiring a FRED API key."""
    url = f"https://fred.stlouisfed.org/graph/fredgraph.csv?id={series_id}"

    try:
        response = requests.get(
            url,
            timeout=HTTP_TIMEOUT,
            headers={"User-Agent": "INFOCHARGE-Market-Intelligence/4.0"},
        )
        response.raise_for_status()

        rows = response.text.splitlines()

        for row in reversed(rows[1:]):
            parts = row.split(",")

            if len(parts) >= 2 and parts[1].strip() not in ("", "."):
                return {
                    "date": parts[0].strip(),
                    "value": float(parts[1].strip()),
                    "series": series_id,
                }

    except Exception as exc:
        logging.warning("FRED series failed: %s | %s", series_id, exc)

    return None


def fred_history(series_id, limit=260):
    """Return recent daily FRED observations for educational charts."""
    url = f"https://fred.stlouisfed.org/graph/fredgraph.csv?id={series_id}"

    try:
        response = requests.get(
            url,
            timeout=HTTP_TIMEOUT,
            headers={"User-Agent": "INFOCHARGE-Market-Intelligence/4.0"},
        )
        response.raise_for_status()

        data = []

        for row in response.text.splitlines()[1:]:
            parts = row.split(",")

            if len(parts) < 2:
                continue

            value = parts[1].strip()

            if value in ("", "."):
                continue

            try:
                data.append((parts[0].strip(), float(value)))
            except ValueError:
                continue

        return data[-limit:]

    except Exception as exc:
        logging.warning("FRED history failed: %s | %s", series_id, exc)
        return []


def fed_balance_sheet_snapshot():
    """
    Reads the current H.4.1 release page and extracts total assets.
    The Federal Reserve releases H.4.1 weekly.
    """
    url = "https://www.federalreserve.gov/releases/h41/current/"

    try:
        response = requests.get(
            url,
            timeout=HTTP_TIMEOUT,
            headers={"User-Agent": "INFOCHARGE-Market-Intelligence/4.0"},
        )
        response.raise_for_status()

        text = clean_text(response.text)

        # Search for Total assets followed by a number in the HTML/text.
        match = re.search(
            r"Total assets\s+([0-9,]+)",
            text,
            flags=re.IGNORECASE,
        )

        if not match:
            # Some pages place the value later in the table.
            match = re.search(
                r"Total assets.*?([0-9,]{6,})",
                text,
                flags=re.IGNORECASE,
            )

        if match:
            return {
                "total_assets_millions": float(
                    match.group(1).replace(",", "")
                ),
                "source": "Federal Reserve H.4.1",
            }

    except Exception as exc:
        logging.warning(
            "Fed balance sheet fetch failed: %s",
            exc
        )

    return None


def macro_market_snapshot():
    return {
        "us_10y": treasury_10y(),
        "us_high_yield_oas": fred_latest("BAMLH0A0HYM2"),
        "fed_balance_sheet": fed_balance_sheet_snapshot(),
        "nifty_it_nasdaq_ratio": nifty_it_nasdaq_ratio(),
    }


def create_nifty_it_nasdaq_ratio_chart():
    """
    Build a historical educational chart from Yahoo daily data.
    No investment signal is implied by the chart.
    """
    if plt is None:
        return None

    it_data = yahoo_chart("^CNXIT", interval="1d")
    if not it_data:
        it_data = yahoo_chart("^NIFTYIT", interval="1d")

    ndx_data = yahoo_chart("^NDX", interval="1d")

    try:
        it_result = it_data["chart"]["result"][0]
        ndx_result = ndx_data["chart"]["result"][0]

        it_ts = it_result["timestamp"]
        it_close = it_result["indicators"]["quote"][0]["close"]

        ndx_ts = ndx_result["timestamp"]
        ndx_close = ndx_result["indicators"]["quote"][0]["close"]

        it_map = {}

        for ts, value in zip(it_ts, it_close):
            if value is not None:
                date_key = datetime.fromtimestamp(
                    ts, tz=IST
                ).strftime("%Y-%m-%d")
                it_map[date_key] = float(value)

        ratio_dates = []
        ratio_values = []

        for ts, value in zip(ndx_ts, ndx_close):
            if value is None:
                continue

            date_key = datetime.fromtimestamp(
                ts, tz=IST
            ).strftime("%Y-%m-%d")

            if date_key in it_map and float(value) != 0:
                ratio_dates.append(date_key)
                ratio_values.append(
                    it_map[date_key] / float(value)
                )

        if len(ratio_values) < 20:
            return None

        # Keep the last ~1 year for a readable educational chart.
        ratio_dates = ratio_dates[-260:]
        ratio_values = ratio_values[-260:]

        fig = plt.figure(figsize=(13, 7))
        ax = fig.add_subplot(111)

        ax.plot(
            ratio_dates,
            ratio_values,
            linewidth=2,
        )

        ax.set_title(
            "Nifty IT / Nasdaq-100 Ratio — Educational Historical View",
            fontsize=16,
            fontweight="bold",
        )
        ax.set_ylabel("Nifty IT ÷ Nasdaq-100")
        ax.set_xlabel("Date")
        ax.grid(alpha=0.25)
        ax.tick_params(axis="x", rotation=45)

        fig.text(
            0.01,
            0.01,
            "Educational only. Ratio is not a standalone buy/sell signal.",
            fontsize=9,
        )

        fig.tight_layout()

        output = Path("nifty_it_nasdaq_ratio.jpg")
        fig.savefig(
            output,
            dpi=150,
            bbox_inches="tight",
        )
        plt.close(fig)

        return output

    except Exception as exc:
        logging.warning(
            "Nifty IT/Nasdaq ratio chart failed: %s",
            exc
        )
        return None


# ============================================================
# CALENDAR / REMINDER ENGINE
# ============================================================

def calendar_event_for_day(now):
    date_key = now.strftime("%Y-%m-%d")
    tomorrow = (now + timedelta(days=1)).strftime("%Y-%m-%d")

    events = []

    if date_key in FOMC_DATES_2026:
        events.append({
            "type": "FOMC",
            "date": date_key,
            "title": "US Federal Reserve / FOMC event",
            "source": FOMC_URL,
        })

    if tomorrow in FOMC_DATES_2026:
        events.append({
            "type": "FOMC_TOMORROW",
            "date": tomorrow,
            "title": "US Federal Reserve / FOMC event tomorrow",
            "source": FOMC_URL,
        })

    if date_key in CPI_RELEASE_DATES_2026:
        events.append({
            "type": "US_CPI",
            "date": date_key,
            "title": "US CPI release",
            "source": CPI_URL,
        })

    if tomorrow in CPI_RELEASE_DATES_2026:
        events.append({
            "type": "US_CPI_TOMORROW",
            "date": tomorrow,
            "title": "US CPI release tomorrow",
            "source": CPI_URL,
        })

    return events


def should_send_calendar(memory, event):
    key = f"{event['type']}:{event['date']}"

    if key in memory.get("calendar_sent", []):
        return False

    return True


def mark_calendar_sent(memory, event):
    key = f"{event['type']}:{event['date']}"
    memory.setdefault("calendar_sent", []).append(key)
    memory["calendar_sent"] = memory["calendar_sent"][-100:]


def calendar_post(event, now):
    prompt = "\n".join([
        "You are INFOCHARGE's macro-event editor.",
        "Write a concise educational Telegram reminder.",
        "Explain why this event can matter for Indian markets without predicting the result.",
        "Mention transmission channels such as USD/INR, US yields, global risk appetite, IT, financials, commodities or FII flows only when relevant.",
        "Do not give a trade call.",
        "Do not invent consensus estimates.",
        "Make it feel like a knowledgeable market person preparing readers.",
        "",
        f"EVENT: {event['title']}",
        f"DATE: {event['date']}",
        f"CURRENT IST: {now.strftime('%d %b %Y, %I:%M %p')}",
        "",
        "Return only the Telegram message."
    ])

    return call_gemini(prompt, 900).strip()


# ============================================================
# CORE EDITORIAL ENGINE
# ============================================================

def research_story(articles, memory, update_type, now):
    macro = macro_market_snapshot()

    context = {
        "macro_market_snapshot": macro,
        "calendar_events": calendar_event_for_day(now),
    }

    prompt = "\n".join([
        "You are the SENIOR MARKET INTELLIGENCE EDITOR for INFOCHARGE INSIDERS CLUB.",
        "Think like a highly knowledgeable human who understands Indian equities, global macro, commodities, geopolitics, corporate filings and sector economics.",
        "",
        f"CURRENT TIME IST: {now.strftime('%d %b %Y, %I:%M %p')}",
        f"UPDATE TYPE: {update_type}",
        f"MARKET STATUS: {market_status(now)}",
        "",
        "MISSION:",
        "Do not behave like an RSS summarizer.",
        "Scan the available information and decide whether there is something genuinely useful to teach or explain.",
        "The channel should be selective. Silence is better than junk.",
        "",
        "WHAT COUNTS AS HIGH-VALUE:",
        "1. A material listed-company order, result, guidance, capex, acquisition, fundraising or strategic change.",
        "2. A sector development with a clear company/earnings/industry transmission mechanism.",
        "3. Geopolitical events ONLY when there is a credible path to Indian stocks/sectors through oil, gas, metals, shipping, trade, tariffs, sanctions, supply chains, defence, currency, rates or other economics.",
        "4. RBI, SEBI, government or regulatory changes with meaningful market/business implications.",
        "5. US Fed/FOMC, CPI, inflation, jobs, Treasury yields or liquidity developments when they materially affect Indian equities or a specific sector.",
        "6. Important IPO developments where investors can learn how to evaluate the issue.",
        "7. An educational market framework using historical data when it genuinely teaches a way of thinking.",
        "8. Rare, high-quality market psychology/motivation content when it adds learning value.",
        "",
        "JUNK FILTER:",
        "Reject generic Nifty/Sensex moves, routine broker commentary, vague optimism/pessimism, tiny orders, promotional company PR, rumours, clickbait, repeated stories, generic geopolitical headlines and news without an Indian market transmission mechanism.",
        "",
        "GEOPOLITICAL TEST:",
        "EVENT -> ECONOMIC CHANNEL -> INDIAN COMPANY/SECTOR EXPOSURE -> POSSIBLE BUSINESS/VALUATION IMPACT -> MATERIALITY.",
        "If one link is weak or speculative, NO_POST.",
        "",
        "EDUCATIONAL TEST:",
        "If using a ratio, yield, historical chart or macro relationship, clearly distinguish FACT from INTERPRETATION.",
        "Never claim correlation proves causation.",
        "Never invent historical numbers.",
        "",
        "IPO TEST:",
        "If the story is an IPO, extract only facts supported by the supplied reports.",
        "Flag missing facts instead of inventing them.",
        "The eventual IPO material must help a reader evaluate the business, growth, margins, cash flow, debt, valuation, issue structure, risks and use of proceeds — not tell them to subscribe.",
        "",
        "DECISION:",
        "Return POST only if the information is genuinely useful to an informed market participant.",
        "Importance 8-10 is required.",
        "",
        "Return ONLY valid JSON with:",
        "decision, content_type, importance, headline, company, sector, key_facts,",
        "market_link, why_it_matters, bigger_theme, caveat, educational_lesson,",
        "memory_key, editorial_note, source_links, ipo_data.",
        "",
        "EDITORIAL MEMORY:",
        memory_for_prompt(memory),
        "",
        "MARKET DATA SNAPSHOTS:",
        json.dumps(context, ensure_ascii=False),
        "",
        "CURRENT REPORTS:",
        article_packet(articles),
    ])

    return extract_json(call_gemini(prompt, 3600))


# ============================================================
# POST WRITER
# ============================================================

def write_post(research, update_type, now):
    prompt = "\n".join([
        "You are the final human editor for INFOCHARGE INSIDERS CLUB.",
        "Write one premium Telegram post from the research below.",
        "",
        "The reader should feel that an experienced market person selected and explained this information.",
        "Do not sound like an RSS bot or AI news digest.",
        "",
        "WRITING RULES:",
        "- Start with the actual insight, not a generic introduction.",
        "- Explain WHAT happened and WHY IT MATTERS.",
        "- If geopolitics: explicitly explain the transmission chain to Indian stocks/sectors.",
        "- If macro: explain what variable matters and which Indian sectors are exposed.",
        "- If educational: teach the framework and how to think, not what to buy.",
        "- If IPO: focus on business quality, numbers, valuation, issue structure and risks.",
        "- Use numbers when supported.",
        "- Short paragraphs; bullets only when useful.",
        "- No buy/sell calls.",
        "- No target prices.",
        "- No guaranteed returns.",
        "- No fearmongering.",
        "- No raw URLs in the prose.",
        "- Do not invent facts.",
        "- Do not force a fixed template.",
        "- Approx 900-1600 characters unless a longer IPO/educational piece is clearly justified.",
        "",
        f"UPDATE TYPE: {update_type}",
        f"TIME: {now.strftime('%d %b %Y, %I:%M %p')}",
        "",
        "RESEARCH:",
        json.dumps(research, ensure_ascii=False, indent=2),
        "",
        "Return ONLY the final Telegram post."
    ])

    return call_gemini(prompt, 2400).strip()


# ============================================================
# EDUCATIONAL CONTENT
# ============================================================

def should_do_education(now, update_type):
    # Education is occasional, not every run.
    # Primarily evening/close/manual education requests.
    return update_type in {"evening", "close", "education"}


def education_post(memory, now):
    macro = macro_market_snapshot()

    prompt = "\n".join([
        "You are INFOCHARGE's market teacher.",
        "Create ONE useful educational market lesson.",
        "The purpose is to teach a way of thinking, not provide a trade.",
        "",
        "Possible subjects:",
        "Fed -> yields -> USD -> India",
        "CPI/inflation -> real rates -> valuation",
        "10Y Treasury yield and equity duration",
        "credit spreads / junk-bond yields as risk appetite",
        "Nifty IT vs Nasdaq relationship",
        "FII flows vs USD/INR",
        "crude oil -> India inflation/current account/sector margins",
        "yield curve and recession thinking",
        "market breadth",
        "earnings revisions",
        "valuation vs growth",
        "how geopolitical shocks transmit to Indian sectors",
        "",
        "Use the available data snapshot if useful.",
        "Never invent historical data.",
        "Clearly label current snapshot versus conceptual explanation.",
        "Do not say a ratio alone predicts markets.",
        "No buy/sell recommendation.",
        "",
        "CURRENT SNAPSHOT:",
        json.dumps(macro, ensure_ascii=False),
        "",
        "RECENT EDUCATION:",
        json.dumps(memory.get("education", [])[-15:], ensure_ascii=False),
        "",
        "Return ONLY the Telegram post."
    ])

    return call_gemini(prompt, 1800).strip()


# ============================================================
# MOTIVATIONAL CONTENT
# ============================================================

def motivation_post(memory, now):
    prompt = "\n".join([
        "Write one short market-related motivational lesson for INFOCHARGE.",
        "It must be about discipline, patience, process, uncertainty, risk, learning or probabilistic thinking.",
        "It should sound original and mature, not like a generic Instagram quote.",
        "Prefer a quote-like line followed by 2-4 sentences explaining the market lesson.",
        "No trading call. No profit promise.",
        "Avoid clichés.",
        "",
        f"DATE: {now.strftime('%d %b %Y')}",
        "Return ONLY the post."
    ])

    return call_gemini(prompt, 900).strip()


# ============================================================
# IPO WATCH
# ============================================================

def ipo_watch_post(research, now):
    prompt = "\n".join([
        "You are an IPO research editor for INFOCHARGE INSIDERS CLUB.",
        "Write an educational IPO due-diligence note from the supplied research.",
        "",
        "Cover, when available:",
        "business model, industry, growth, revenue, EBITDA/PAT, margins, cash flow, debt,",
        "issue size, fresh issue vs OFS, price/valuation, use of proceeds, promoters,",
        "competitive position, key risks, important positives, and what information is still missing.",
        "",
        "Do NOT say subscribe/avoid.",
        "Do NOT give a target price.",
        "Do NOT invent issue dates, price bands, lot size, valuation or financial numbers.",
        "If a fact is missing, explicitly say that it needs verification from the RHP/DRHP/official filing.",
        "Tone: analytical and balanced.",
        "This is an educational framework, not investment advice.",
        "",
        "RESEARCH:",
        json.dumps(research, ensure_ascii=False, indent=2),
        "",
        f"TIME: {now.strftime('%d %b %Y, %I:%M %p')}",
        "Return ONLY the Telegram post."
    ])

    return call_gemini(prompt, 3000).strip()


# ============================================================
# FREE VISUAL CARD
# ============================================================

def create_visual_card(title, subtitle="", footer="INFOCHARGE INSIDERS CLUB"):
    """
    Creates a simple branded image locally with Pillow.
    No paid image API and no external image copyright issue.
    """
    try:
        from PIL import Image, ImageDraw, ImageFont

        width, height = 1600, 900
        image = Image.new("RGB", (width, height), (239, 247, 255))
        draw = ImageDraw.Draw(image)

        # Find common fonts; fall back safely.
        font_paths = [
            "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf",
            "/usr/share/fonts/truetype/liberation2/LiberationSans-Bold.ttf",
        ]
        regular_paths = [
            "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
            "/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf",
        ]

        bold_path = next((p for p in font_paths if Path(p).exists()), None)
        regular_path = next((p for p in regular_paths if Path(p).exists()), None)

        if bold_path:
            big = ImageFont.truetype(bold_path, 78)
            medium = ImageFont.truetype(bold_path, 44)
        else:
            big = ImageFont.load_default()
            medium = ImageFont.load_default()

        if regular_path:
            small = ImageFont.truetype(regular_path, 30)
        else:
            small = ImageFont.load_default()

        # Editorial visual, not a fake photograph.
        draw.rectangle((0, 0, width, 18), fill=(20, 90, 150))
        draw.text((90, 80), "INFOCHARGE", font=medium, fill=(20, 70, 120))

        # Wrap title.
        words = title.split()
        lines = []
        line = ""

        for word in words:
            test = f"{line} {word}".strip()

            if draw.textlength(test, font=big) <= 1400:
                line = test
            else:
                if line:
                    lines.append(line)
                line = word

        if line:
            lines.append(line)

        y = 220

        for line in lines[:4]:
            draw.text((90, y), line, font=big, fill=(15, 25, 35))
            y += 100

        if subtitle:
            subtitle = clean_text(subtitle)

            words = subtitle.split()
            lines2 = []
            line = ""

            for word in words:
                test = f"{line} {word}".strip()

                if draw.textlength(test, font=small) <= 1380:
                    line = test
                else:
                    if line:
                        lines2.append(line)
                    line = word

            if line:
                lines2.append(line)

            y += 30

            for line in lines2[:3]:
                draw.text((90, y), line, font=small, fill=(60, 75, 90))
                y += 48

        draw.text(
            (90, 810),
            footer,
            font=small,
            fill=(20, 90, 150),
        )

        output = Path("infocharge_visual.jpg")
        image.save(output, quality=92, optimize=True)

        return output

    except Exception as exc:
        logging.warning("Visual card creation failed: %s", exc)
        return None



def create_ipo_ppt(research, post):
    """
    Create a concise IPO due-diligence deck from verified research fields.
    It never fabricates missing numbers.
    """
    try:
        from pptx import Presentation
        from pptx.util import Inches, Pt

        company = research.get("company") or "IPO Watch"
        headline = research.get("headline") or company
        data = research.get("ipo_data") or {}

        prs = Presentation()
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.5)

        def add_slide(title, body):
            slide = prs.slides.add_slide(prs.slide_layouts[5])
            title_box = slide.shapes.add_textbox(
                Inches(0.6), Inches(0.35), Inches(12.1), Inches(0.8)
            )
            title_tf = title_box.text_frame
            title_tf.text = title
            title_tf.paragraphs[0].font.size = Pt(28)
            title_tf.paragraphs[0].font.bold = True

            body_box = slide.shapes.add_textbox(
                Inches(0.7), Inches(1.35), Inches(11.9), Inches(5.6)
            )
            tf = body_box.text_frame
            tf.word_wrap = True
            tf.text = body

            for paragraph in tf.paragraphs:
                paragraph.font.size = Pt(17)

            return slide

        add_slide(
            f"{company} — IPO Intelligence",
            f"{headline}\n\nEducational research prepared by INFOCHARGE INSIDERS CLUB.\n"
            "Not investment advice. Verify all issue details from the latest DRHP/RHP and exchange filings."
        )

        facts = research.get("key_facts") or []
        add_slide(
            "1. What is the business?",
            "\n".join(f"• {x}" for x in facts[:8]) or
            "Verified business facts were not sufficiently available in the supplied reports."
        )

        add_slide(
            "2. Issue structure & valuation",
            json.dumps(
                data,
                ensure_ascii=False,
                indent=2
            )[:7000] or
            "Issue size, price band, OFS/fresh issue, lot size and valuation: verify from the latest RHP/official filing."
        )

        add_slide(
            "3. Why the business may be interesting",
            research.get("why_it_matters")
            or "Not enough verified information."
        )

        add_slide(
            "4. Risks / what can go wrong",
            research.get("caveat")
            or "Risk analysis requires verification from the latest offer document."
        )

        add_slide(
            "5. How to think about it",
            research.get("educational_lesson")
            or "Compare growth, margins, cash flow, debt, valuation, competition and governance before forming a view."
        )

        add_slide(
            "6. INFOCHARGE conclusion framework",
            "Do not ask: “Will this IPO list at a premium?”\n\n"
            "Ask:\n"
            "• What am I paying?\n"
            "• What growth am I underwriting?\n"
            "• Are cash flows real?\n"
            "• What are the key risks?\n"
            "• What changes if growth disappoints?\n\n"
            "This deck is educational and is not a subscribe/avoid recommendation."
        )

        output = Path("INFOCHARGE_IPO_WATCH.pptx")
        prs.save(output)
        return output

    except Exception as exc:
        logging.warning("IPO PPT creation failed: %s", exc)
        return None


def send_document(path, caption=""):
    with open(path, "rb") as document:
        telegram_request(
            "sendDocument",
            payload={
                "chat_id": TELEGRAM_CHAT_ID,
                "caption": caption[:1000],
            },
            files={
                "document": document,
            },
        )


# ============================================================
# TELEGRAM
# ============================================================

def split_message(text, limit=3900):
    if len(text) <= limit:
        return [text]

    chunks = []
    current = ""

    for paragraph in text.split("\n\n"):
        candidate = paragraph if not current else current + "\n\n" + paragraph

        if len(candidate) <= limit:
            current = candidate
        else:
            if current:
                chunks.append(current)
            current = paragraph

    if current:
        chunks.append(current)

    result = []

    for chunk in chunks:
        if len(chunk) <= limit:
            result.append(chunk)
        else:
            result.extend(
                chunk[i:i + limit]
                for i in range(0, len(chunk), limit)
            )

    return result


def telegram_request(method, payload=None, files=None):
    url = f"https://api.telegram.org/bot{TELEGRAM_BOT_TOKEN}/{method}"

    response = requests.post(
        url,
        data=payload,
        files=files,
        timeout=HTTP_TIMEOUT,
    )

    response.raise_for_status()

    result = response.json()

    if not result.get("ok"):
        raise RuntimeError(f"Telegram API error: {result}")

    return result


def send_text(text):
    for chunk in split_message(text):
        telegram_request(
            "sendMessage",
            payload={
                "chat_id": TELEGRAM_CHAT_ID,
                "text": chunk,
                "disable_web_page_preview": "true",
            },
        )


def send_photo(photo_path, caption):
    with open(photo_path, "rb") as image_file:
        telegram_request(
            "sendPhoto",
            payload={
                "chat_id": TELEGRAM_CHAT_ID,
                "caption": caption[:1000],
            },
            files={
                "photo": image_file,
            },
        )


def send_post_with_optional_visual(post, research):
    """
    Use a visual card for high-value posts.
    If card creation fails, safely send text.
    """
    title = research.get("headline", "INFOCHARGE MARKET INTELLIGENCE")
    subtitle = research.get("why_it_matters", "")

    visual = create_visual_card(title, subtitle)

    if visual and visual.exists():
        try:
            send_photo(visual, post)
            return
        except Exception as exc:
            logging.warning(
                "Photo send failed; falling back to text: %s",
                exc
            )

    send_text(post)


# ============================================================
# MEMORY
# ============================================================

def record_memory(memory, research, post, content_type="news"):
    memory.setdefault("posted", []).append({
        "date": datetime.now(IST).strftime("%Y-%m-%d %H:%M"),
        "headline": research.get("headline", ""),
        "company": research.get("company", ""),
        "key": research.get("memory_key", ""),
        "importance": research.get("importance", 0),
        "content_type": content_type,
        "post": post,
    })

    memory["posted"] = memory["posted"][-MAX_MEMORY_ITEMS:]

    note = research.get("editorial_note")

    if note:
        memory.setdefault("editorial_notes", []).append(note)
        memory["editorial_notes"] = memory["editorial_notes"][-30:]


# ============================================================
# MAIN
# ============================================================

def main():
    now = datetime.now(IST)
    update_type = os.getenv("INFOCHARGE_UPDATE_TYPE", "hourly").lower()

    logging.info(
        "INFOCHARGE V4 | %s | update=%s | status=%s",
        now.isoformat(),
        update_type,
        market_status(now),
    )

    memory = load_memory()

    # --------------------------------------------------------
    # 1. Calendar reminder layer
    # --------------------------------------------------------
    if update_type in {"premarket", "open", "hourly", "close", "evening", "macro"}:
        events = calendar_event_for_day(now)

        for event in events:
            if should_send_calendar(memory, event):
                try:
                    message = calendar_post(event, now)
                    send_text(message)
                    mark_calendar_sent(memory, event)
                    save_memory(memory)

                    logging.info(
                        "Calendar reminder sent: %s",
                        event["type"]
                    )
                except Exception as exc:
                    logging.warning(
                        "Calendar reminder failed: %s",
                        exc
                    )

    # --------------------------------------------------------
    # 2. Market closed: education/calendar only
    # --------------------------------------------------------
    if market_status(now) == "closed":
        if update_type in {"evening", "education"}:
            try:
                post = education_post(memory, now)
                send_text(post)

                memory.setdefault("education", []).append({
                    "date": now.strftime("%Y-%m-%d"),
                    "post": post,
                })
                memory["education"] = memory["education"][-30:]
                save_memory(memory)

                logging.info("Closed-day education post sent.")
            except Exception as exc:
                logging.warning(
                    "Closed-day education failed: %s",
                    exc
                )

        return

    # --------------------------------------------------------
    # 3. Education mode
    # --------------------------------------------------------
    if update_type == "education":
        post = education_post(memory, now)

        chart = create_nifty_it_nasdaq_ratio_chart()

        if chart and chart.exists():
            try:
                send_photo(chart, post)
            except Exception as exc:
                logging.warning(
                    "Educational chart send failed: %s",
                    exc
                )
                send_text(post)
        else:
            send_text(post)

        memory.setdefault("education", []).append({
            "date": now.strftime("%Y-%m-%d"),
            "post": post,
        })
        memory["education"] = memory["education"][-30:]
        save_memory(memory)

        logging.info("Education post sent.")
        return

    # --------------------------------------------------------
    # 4. Occasional motivation mode
    # --------------------------------------------------------
    if update_type == "motivation":
        post = motivation_post(memory, now)
        send_text(post)
        logging.info("Motivation post sent.")
        return

    # --------------------------------------------------------
    # 5. Normal intelligence scan
    # --------------------------------------------------------
    articles = fetch_news()

    logging.info(
        "Collected %d unique reports.",
        len(articles)
    )

    if not articles:
        logging.info("No reports collected. Nothing to post.")
        return

    research = research_story(
        articles,
        memory,
        update_type,
        now,
    )

    decision = str(
        research.get("decision", "NO_POST")
    ).upper()

    importance = safe_int(
        research.get("importance", 0),
        0,
    )

    content_type = str(
        research.get("content_type", "news")
    ).lower()

    logging.info(
        "Decision=%s | importance=%s | type=%s | headline=%s",
        decision,
        importance,
        content_type,
        research.get("headline", ""),
    )

    if decision != "POST" or importance < MIN_IMPORTANCE:
        logging.info(
            "Nothing passed the high-value editorial filter."
        )
        return

    # --------------------------------------------------------
    # 6. IPO-specific writer
    # --------------------------------------------------------
    if content_type == "ipo":
        post = ipo_watch_post(research, now)
    else:
        post = write_post(
            research,
            update_type,
            now,
        )

    if not post:
        raise RuntimeError("Final post was empty.")

    # --------------------------------------------------------
    # 7. Send with free visual card
    # --------------------------------------------------------
    send_post_with_optional_visual(
        post,
        research,
    )

    # IPO mode gets a structured research deck as well.
    if content_type == "ipo":
        ppt = create_ipo_ppt(research, post)

        if ppt and ppt.exists():
            try:
                send_document(
                    ppt,
                    caption=f"INFOCHARGE IPO Intelligence — {research.get('company', 'IPO Watch')}"
                )
            except Exception as exc:
                logging.warning(
                    "IPO PPT send failed: %s",
                    exc
                )

    # --------------------------------------------------------
    # 8. Persistent memory
    # --------------------------------------------------------
    record_memory(
        memory,
        research,
        post,
        content_type=content_type,
    )

    if content_type == "ipo":
        memory.setdefault("ipo_watch", []).append({
            "date": now.strftime("%Y-%m-%d"),
            "company": research.get("company", ""),
            "headline": research.get("headline", ""),
            "data": research.get("ipo_data", {}),
        })
        memory["ipo_watch"] = memory["ipo_watch"][-30:]

    save_memory(memory)

    logging.info(
        "INFOCHARGE V4 post sent and memory saved."
    )


if __name__ == "__main__":
    main()
